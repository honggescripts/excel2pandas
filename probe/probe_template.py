# -*- coding: utf-8 -*-
"""
模板探针 —— 接手一个新 Excel 模板时，先用它摸清结构，再决定怎么迁移。

做五件事
--------
1. sheet 清单与规模
2. 目标 sheet 的真实行列数
   ⚠️ openpyxl 只读模式下 `<dimension>` 元数据不可信，必须先 reset_dimensions()；
      而 reset 之后 ws.max_row / ws.max_column 会变成 None，所以规模只能靠扫描累计。
3. 函数直方图 —— 决定本期要支持哪些 Excel 函数
4. 跨表引用清单 + 外部工作簿引用 —— 判断向量化可行性与缺口
5. 绝对引用（$）/ 区域引用（A1:B1）统计 —— 预判解析难点

用法
----
    python probe/probe_template.py --file demo.xlsx --sheet 定价测算 --row 2
    python probe/probe_template.py --file demo.xlsx --sheet 定价测算 --row 2 --json out.json

产物
----
默认只打印摘要；给了 --json 才落盘（内容含表头与样板行原文，注意不要提交到公开仓库）。
"""

import argparse
import json
import re
from collections import Counter
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

FUNC_RE = re.compile(r"([A-Za-z][A-Za-z0-9\.]*)\s*\(")
RANGE_RE = re.compile(r"\$?[A-Z]{1,3}\$?[0-9]+\s*:\s*\$?[A-Z]{1,3}\$?[0-9]+")
ABSREF_RE = re.compile(r"\$[A-Z]{1,3}\$?[0-9]+|\$?[A-Z]{1,3}\$[0-9]+")
SHEETREF_RE = re.compile(r"(?:'([^']+)'|([A-Za-z0-9\u4e00-\u9fa5_\.]+))!")
EXTBOOK_RE = re.compile(r"\[([^\]]+)\]")


def probe(xlsx: Path, sheet: str, sample_row: int, max_rows: int, sample_cols: int) -> dict:
    res = {"file": xlsx.name, "size_mb": round(xlsx.stat().st_size / 1024 / 1024, 2)}

    wb = openpyxl.load_workbook(str(xlsx), read_only=True, data_only=False, keep_links=False)
    res["sheetnames"] = list(wb.sheetnames)

    if sheet not in wb.sheetnames:
        wb.close()
        raise SystemExit(f"sheet 不存在：{sheet!r}\n可选：{res['sheetnames']}")

    ws = wb[sheet]
    ws.reset_dimensions()          # 关键：丢弃不可信的 <dimension>，强制全表扫描

    func_counter, sheetref_counter = Counter(), Counter()
    extbooks = set()
    formula_cells = value_cells = 0
    range_hits = abs_hits = 0
    rows_with_formula = 0
    max_r_seen = max_c_seen = 0
    header, sample = {}, {}

    for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=max_rows), start=1):
        max_r_seen = r_idx
        f_in_row = 0
        for c_idx, cell in enumerate(row, start=1):
            if c_idx > max_c_seen:
                max_c_seen = c_idx
            val = getattr(cell, "value", None)
            if val is None:
                continue
            col = get_column_letter(c_idx)
            if isinstance(val, str) and val.startswith("="):
                f_in_row += 1
                formula_cells += 1
                body = val[1:]
                for m in FUNC_RE.finditer(body):
                    func_counter[m.group(1).upper()] += 1
                for m in SHEETREF_RE.finditer(body):
                    sheetref_counter[m.group(1) or m.group(2)] += 1
                for m in EXTBOOK_RE.finditer(body):
                    extbooks.add(m.group(1))
                if RANGE_RE.search(body):
                    range_hits += 1
                if ABSREF_RE.search(body):
                    abs_hits += 1
            else:
                value_cells += 1
        if f_in_row:
            rows_with_formula += 1
        if r_idx == 1:
            header = {
                get_column_letter(i): (str(c.value)[:30] if c.value is not None else "")
                for i, c in enumerate(row, start=1) if i <= sample_cols
            }
        if r_idx == sample_row:
            sample = {
                get_column_letter(i): (str(c.value)[:80] if c.value is not None else "")
                for i, c in enumerate(row, start=1) if c.value is not None and i <= sample_cols
            }
    wb.close()

    res["dims"] = {"rows": max_r_seen, "cols": max_c_seen}
    if max_r_seen >= max_rows:
        res["dims"]["truncated"] = f"已扫描到 --max-rows={max_rows} 上限，实际可能更大"
    res["formula_cells"] = formula_cells
    res["value_cells"] = value_cells
    res["rows_with_formula"] = rows_with_formula
    res["ref_flags"] = {"含绝对引用($)": abs_hits, "含区域引用(A1:B1)": range_hits}
    res["func_hist"] = dict(func_counter.most_common(60))
    res["sheetref_hist"] = dict(sheetref_counter.most_common(40))
    res["external_books"] = sorted(extbooks)
    res["header"] = header
    res["sample_row"] = sample
    return res


def main():
    ap = argparse.ArgumentParser(description="Excel 模板结构探针")
    ap.add_argument("--file", required=True, help="xlsx 路径")
    ap.add_argument("--sheet", required=True, help="目标 sheet 名")
    ap.add_argument("--row", type=int, default=2, help="样板行号（默认 2）")
    ap.add_argument("--max-rows", type=int, default=5000, help="安全扫描上限")
    ap.add_argument("--sample-cols", type=int, default=120)
    ap.add_argument("--json", default=None, help="把完整结果写入指定文件")
    a = ap.parse_args()

    xlsx = Path(a.file)
    if not xlsx.exists():
        raise SystemExit(f"文件不存在：{xlsx}")

    res = probe(xlsx, a.sheet, a.row, a.max_rows, a.sample_cols)

    print(f"[1] {res['file']}  ({res['size_mb']} MB)")
    print(f"[2] sheet 共 {len(res['sheetnames'])} 个：{res['sheetnames'][:8]}"
          f"{' …' if len(res['sheetnames']) > 8 else ''}")
    print(f"[3] 目标 sheet={a.sheet}  {res['dims']['rows']} 行 × {res['dims']['cols']} 列")
    if res["dims"].get("truncated"):
        print(f"    ⚠️ {res['dims']['truncated']}")
    print(f"    公式单元格={res['formula_cells']}  值单元格={res['value_cells']}  "
          f"含公式行数={res['rows_with_formula']}")
    print(f"    含绝对引用={res['ref_flags']['含绝对引用($)']}  "
          f"含区域引用={res['ref_flags']['含区域引用(A1:B1)']}")

    print("\n[4] 函数直方图 TOP25（决定本期支持范围）")
    for k, v in list(res["func_hist"].items())[:25]:
        print(f"      {k:<16} {v}")

    print("\n[5] 跨表引用 TOP20")
    for k, v in list(res["sheetref_hist"].items())[:20]:
        print(f"      {k:<28} {v}")
    print(f"    跨工作簿引用：{res['external_books'] or '无'}")

    if a.json:
        Path(a.json).write_text(json.dumps(res, ensure_ascii=False, indent=1), encoding="utf-8")
        print(f"\n-> {a.json}  （含表头/样板行原文，勿提交到公开仓库）")


if __name__ == "__main__":
    main()

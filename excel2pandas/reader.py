# -*- coding: utf-8 -*-
"""
reader.py —— 读簿建模型（合并原设计的 reader + layer + assessor）

职责：
  1. 读目标 sheet 的表头行 / 样板行（双读：公式 + 缓存值）
  2. 列建模与列名唯一化
  3. 3+1 层分类 + 分层冲突检测
  4. 迁移评估报告（migration_report.md）

不做：公式转换、代码生成、执行。
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

# ---------------------------------------------------------------- 常量

EXCEL_ERRORS = {"#REF!", "#N/A", "#DIV/0!", "#VALUE!", "#NUM!", "#NAME?", "#NULL!", "#SPILL!", "#CALC!"}

# Layer 2 的判据函数（公式里出现即归 Layer 2）
LOOKUP_FUNCS = {"VLOOKUP", "XLOOKUP", "HLOOKUP", "INDEX"}
AGG_FUNCS = {"SUMIFS", "COUNTIFS", "AVERAGEIFS", "MAXIFS", "MINIFS", "SUMIF", "COUNTIF", "AVERAGEIF"}

HEADER_ROW = 1
DEFAULT_TARGET_ROW = 2

_REF_ONLY_RE = re.compile(
    r"^=\s*(?:'(?P<q>[^']+)'|(?P<p>[A-Za-z0-9\u4e00-\u9fa5_\.]+))?\s*!?\s*"
    r"(?P<c>\$?[A-Z]{1,3})(?P<r>\$?\d+)\s*$"
)
_SHEET_PREFIX_RE = re.compile(r"(?:'([^']+)'|([A-Za-z0-9\u4e00-\u9fa5_\.]+))!")


def norm_header(v: Any) -> str:
    """表头归一化：折行/多空格压平。"""
    if v is None:
        return ""
    s = re.sub(r"\s+", " ", str(v)).strip()
    return s


def is_excel_error(v: Any) -> bool:
    return isinstance(v, str) and v.strip().upper() in EXCEL_ERRORS


def parse_numeric_text(v: Any) -> Optional[float]:
    """'12.5%' -> 0.125、'1,280' -> 1280.0；不是数字文本则 None。

    Excel 在算术运算里会隐式把这类文本转成数字，
    生成代码必须复现这个行为，否则 '12.5%'/12 会直接抛 TypeError。
    """
    if not isinstance(v, str):
        return None
    s = v.strip()
    if not s:
        return None
    if s.endswith("%"):
        try:
            return float(s[:-1].replace(",", "")) / 100.0
        except ValueError:
            return None
    try:
        return float(s.replace(",", ""))
    except ValueError:
        return None


def col_letter(ref: str) -> str:
    return re.match(r"\$?([A-Z]{1,3})", ref).group(1)


# ---------------------------------------------------------------- 数据模型


@dataclass
class ColumnMeta:
    letter: str
    index: int
    header: str
    name: str                      # 唯一化后的输出列名
    duplicated: bool = False
    kind: str = "empty"            # data / formula / empty
    layer: str = "data"            # data / L1 / L2 / L3
    formula: Optional[str] = None
    cached: Any = None
    cached_is_error: bool = False
    cached_missing: bool = False
    cached_numeric_text: bool = False      # 缓存值是「数字样式的文本」-> 参与运算需 _n() 转换
    depends_on: Set[str] = field(default_factory=set)      # 同表列字母
    same_ranges: List[Tuple[str, str]] = field(default_factory=list)
    ref_sheets: Set[str] = field(default_factory=set)      # 引用到的外部 sheet
    ref_ext_cols: Dict[str, Set[str]] = field(default_factory=dict)
    ref_ext_ranges: Dict[str, List[Tuple[str, str]]] = field(default_factory=dict)
    unsupported: List[str] = field(default_factory=list)   # 遇到的未支持函数


@dataclass
class ExternalSheet:
    name: str
    letters: Dict[str, str] = field(default_factory=dict)   # 列字母 -> 表头
    n_rows: Optional[int] = None                            # None = 未扫描


@dataclass
class SheetModel:
    path: str
    sheet_name: str
    header_row: int
    target_row: int
    primary_key_letter: Optional[str]
    columns: List[ColumnMeta]
    externals: Dict[str, ExternalSheet]
    conflicts: List[str] = field(default_factory=list)

    # ---- 便捷索引 ----
    @property
    def by_letter(self) -> Dict[str, ColumnMeta]:
        return {c.letter: c for c in self.columns}

    @property
    def by_name(self) -> Dict[str, ColumnMeta]:
        return {c.name: c for c in self.columns}

    def pick(self, letter: str) -> Optional[ColumnMeta]:
        return self.by_letter.get(letter)

    def resolve_name(self, letter: str) -> Optional[str]:
        c = self.by_letter.get(letter)
        return c.name if c else None

    @property
    def formula_cols(self) -> List[ColumnMeta]:
        return [c for c in self.columns if c.kind == "formula"]

    @property
    def data_cols(self) -> List[ColumnMeta]:
        return [c for c in self.columns if c.kind == "data"]

    @property
    def layer1(self) -> List[ColumnMeta]:
        return [c for c in self.columns if c.layer == "L1"]

    @property
    def layer2(self) -> List[ColumnMeta]:
        return [c for c in self.columns if c.layer == "L2"]

    @property
    def layer3(self) -> List[ColumnMeta]:
        return [c for c in self.columns if c.layer == "L3"]

    @property
    def fail_cols(self) -> List[ColumnMeta]:
        return [c for c in self.formula_cols if c.unsupported]


# ---------------------------------------------------------------- 分层


def classify_layer(formula: Optional[str]) -> str:
    """公式 -> data / L1 / L2 / L3"""
    if not formula:
        return "data"
    body = formula.lstrip("=").strip()
    if _REF_ONLY_RE.match("=" + body):
        return "L1"
    upper = body.upper()
    names = set(re.findall(r"([A-Za-z_][A-Za-z0-9_\.]*)\s*\(", upper))
    names = {n.split(".")[-1] for n in names}
    if names & LOOKUP_FUNCS:
        if "INDEX" in names and "MATCH" not in names:
            pass  # 单独的 INDEX 不算 lookup
        else:
            return "L2"
    if names & AGG_FUNCS:
        return "L2"
    return "L3"


@dataclass
class RefInfo:
    """一条公式里解析出来的全部引用。"""
    same_cols: Set[str] = field(default_factory=set)                    # 同表：列字母
    same_ranges: List[Tuple[str, str]] = field(default_factory=list)    # 同表：区域 (起列, 止列)
    ext_cols: Dict[str, Set[str]] = field(default_factory=dict)         # 跨表：sheet -> 列字母
    ext_ranges: Dict[str, List[Tuple[str, str]]] = field(default_factory=dict)
    has_error_ref: bool = False


def _ref_cols_of(token: str) -> List[str]:
    """从 'B:B' / 'B2:H2' / 'B2' 里抠出涉及的列字母。"""
    return re.findall(r"\$?([A-Z]{1,3})(?=\$?\d|\s*:|\s*$)", token)


def _slice_range(token: str) -> Tuple[str, str]:
    """'B2:H2' -> ('B','H')"""
    parts = token.split(":")
    if len(parts) == 2:
        return col_letter(parts[0]), col_letter(parts[1])
    c = col_letter(token)
    return c, c


# 「可选sheet前缀 + ! + 引用」—— 跨表引用必须先整体遮蔽，否则整列引用会丢掉归属
_QUALIFIED_RE = re.compile(
    r"(?:'(?P<q>[^']+)'|(?P<p>[A-Za-z0-9\u4e00-\u9fa5_\.]+))"
    r"\s*!\s*"
    r"(?P<ref>\$?[A-Z]{1,3}\$?\d+\s*:\s*\$?[A-Z]{1,3}\$?\d+"     # A1:B2
    r"|\$?[A-Z]{1,3}\s*:\s*\$?[A-Z]{1,3}(?![A-Za-z0-9_])"        # A:B
    r"|\$?[A-Z]{1,3}\$?\d+(?![A-Za-z0-9_]))"                     # A1
)

# 同表引用（剩下的）
_SAME_CELL_RE = re.compile(r"(?<![A-Za-z0-9_\$\.!])\$?([A-Z]{1,3})\$?\d+(?![A-Za-z0-9_])")
_SAME_RANGE_RE = re.compile(
    r"(?<![A-Za-z0-9_\$\.!])\$?([A-Z]{1,3})\$?\d+\s*:\s*\$?([A-Z]{1,3})\$?\d+(?![A-Za-z0-9_])"
    r"|(?<![A-Za-z0-9_\$\.!])\$?([A-Z]{1,3})\s*:\s*\$?([A-Z]{1,3})(?![A-Za-z0-9_])"
)


def extract_refs(formula: str, sheet: str, known_sheets: Set[str]) -> RefInfo:
    """
    解析公式里所有引用。
    步骤：① 先把带 sheet 前缀的引用整体遮蔽并登记；② 再在剩余文本里找同表引用。
    这样 `'运费价目'!B:B` 不会被误当成同表 B 列。
    """
    out = RefInfo()
    body = formula.lstrip("=")

    def _mask(m: "re.Match[str]") -> str:
        sname = m.group("q") or m.group("p")
        token = m.group("ref")
        is_range = ":" in token
        if sname in known_sheets and sname != sheet:
            if is_range:
                out.ext_ranges.setdefault(sname, []).append(_slice_range(token))
                for c in _ref_cols_of(token):
                    out.ext_cols.setdefault(sname, set()).add(c)
            else:
                for c in _ref_cols_of(token):
                    out.ext_cols.setdefault(sname, set()).add(c)
        else:
            # `本表名!A1` 这种写法也算同表引用
            if is_range:
                out.same_ranges.append(_slice_range(token))
            for c in _ref_cols_of(token):
                out.same_cols.add(c)
        return " " * (m.end() - m.start())

    masked = _QUALIFIED_RE.sub(_mask, body)

    for m in _SAME_RANGE_RE.finditer(masked):
        a = m.group(1) or m.group(3)
        b = m.group(2) or m.group(4)
        out.same_ranges.append((a, b))
        out.same_cols.update({a, b})
    for m in _SAME_CELL_RE.finditer(masked):
        out.same_cols.add(m.group(1))

    return out



# ---------------------------------------------------------------- 主入口


def read_model(
    path: str | Path,
    sheet_name: str,
    target_row: int = DEFAULT_TARGET_ROW,
    header_row: int = HEADER_ROW,
    max_col: Optional[int] = None,
    scan_external: bool = False,
) -> SheetModel:
    path = str(Path(path).resolve())

    wb_f = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
    if sheet_name not in wb_f.sheetnames:
        raise KeyError(f"未找到 sheet：{sheet_name}（可选：{wb_f.sheetnames}）")
    known_sheets = set(wb_f.sheetnames)

    ws_f = wb_f[sheet_name]
    ws_f.reset_dimensions()
    rows_f = list(ws_f.iter_rows(min_row=header_row, max_row=target_row))
    wb_f.close()

    wb_v = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
    ws_v = wb_v[sheet_name]
    ws_v.reset_dimensions()
    rows_v = list(ws_v.iter_rows(min_row=header_row, max_row=target_row))
    wb_v.close()

    hdr_f = rows_f[0] if rows_f else []
    hdr_v = rows_v[0] if rows_v else []
    tgt_f = rows_f[-1] if len(rows_f) > 1 else []
    tgt_v = rows_v[-1] if len(rows_v) > 1 else []

    n = max([len(hdr_f), len(hdr_v), len(tgt_f), len(tgt_v), max_col or 0])

    # ---------- 列建模 ----------
    columns: List[ColumnMeta] = []
    seen: Dict[str, int] = {}
    for i in range(1, n + 1):
        letter = get_column_letter(i)
        header = norm_header(hdr_f[i - 1].value if i <= len(hdr_f) else None)
        if not header:
            header = norm_header(hdr_v[i - 1].value if i <= len(hdr_v) else None)

        f_cell = tgt_f[i - 1] if i <= len(tgt_f) else None
        v_cell = tgt_v[i - 1] if i <= len(tgt_v) else None
        fval = getattr(f_cell, "value", None)
        vval = getattr(v_cell, "value", None)

        formula = fval if isinstance(fval, str) and fval.startswith("=") else None
        if formula:
            kind = "formula"
        elif vval is not None or fval is not None:
            kind = "data"
        else:
            kind = "empty"

        # 输出列名唯一化
        base = header or letter
        if base in seen:
            seen[base] += 1
            name = f"{base}_{seen[base]}"
            dup = True
        else:
            seen[base] = 1
            name = base
            dup = False

        columns.append(
            ColumnMeta(
                letter=letter,
                index=i,
                header=header,
                name=name,
                duplicated=dup,
                kind=kind,
                formula=formula,
                cached=vval,
                cached_is_error=is_excel_error(vval),
                cached_missing=(vval is None and formula is not None),
                cached_numeric_text=(parse_numeric_text(vval) is not None),
            )
        )

    # 去掉尾部连续空列
    while columns and columns[-1].kind == "empty" and not columns[-1].header:
        columns.pop()
    n = len(columns)

    # ---------- 分层 + 依赖 ----------
    for c in columns:
        c.layer = classify_layer(c.formula)
    letters = {c.letter for c in columns}
    for c in columns:
        if not c.formula:
            continue
        info = extract_refs(c.formula, sheet_name, known_sheets)
        c.depends_on = {x for x in info.same_cols if x != c.letter and x in letters}
        c.same_ranges = info.same_ranges
        c.ref_sheets = set(info.ext_cols) | set(info.ext_ranges)
        c.ref_ext_cols = info.ext_cols
        c.ref_ext_ranges = info.ext_ranges

    # ---------- 主键列推断：优先「主键 / ID / 编号 / 序号」，否则取第一个非公式列 ----------
    pk = None
    for key in ("主键", "ID", "id", "编号", "KEY", "序号"):
        for c in columns:
            if c.header.replace(" ", "") in (key,) or c.header.startswith(key):
                if c.kind == "data":
                    pk = c.letter
                    break
        if pk:
            break
    if pk is None:
        for c in columns:
            if c.kind == "data":
                pk = c.letter
                break

    # ---------- 分层冲突检测 ----------
    # 定义：分层模型要求「低层不依赖高层」。出现 L1 -> L2/L3 或 L2 -> L3 的边即为冲突，
    #       此时走全局拓扑降级。
    conflicts: List[str] = []
    rank = {"data": 0, "L1": 1, "L2": 2, "L3": 3}
    for c in columns:
        if not c.formula:
            continue
        for d in c.depends_on:
            dc = next((x for x in columns if x.letter == d), None)
            if dc and rank[dc.layer] > rank[c.layer]:
                conflicts.append(
                    f"{c.letter}({c.name}·{c.layer}) 引用了更高层的 {dc.letter}({dc.name}·{dc.layer})"
                )

    # ---------- 外部 sheet 表头 ----------
    externals: Dict[str, ExternalSheet] = {}
    need = sorted({s for c in columns for s in c.ref_sheets})
    if need:
        wb_e = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
        for name in need:
            if name not in wb_e.sheetnames:
                continue
            ws_e = wb_e[name]
            ws_e.reset_dimensions()
            letters: Dict[str, str] = {}
            hrow = next(ws_e.iter_rows(min_row=1, max_row=1), [])
            for i, cell in enumerate(hrow, start=1):
                letters[get_column_letter(i)] = norm_header(getattr(cell, "value", None))
            es = ExternalSheet(name=name, letters=letters)
            if scan_external:
                cnt = 0
                for cnt, _ in enumerate(ws_e.iter_rows(min_row=1), start=1):
                    pass
                es.n_rows = cnt
            externals[name] = es
        wb_e.close()

    return SheetModel(
        path=path,
        sheet_name=sheet_name,
        header_row=header_row,
        target_row=target_row,
        primary_key_letter=pk,
        columns=columns,
        externals=externals,
        conflicts=conflicts,
    )


# ---------------------------------------------------------------- 报告


def write_migration_report(model: SheetModel, out_path: str | Path) -> str:
    cols = model.columns
    f_cols = model.formula_cols
    d_cols = model.data_cols
    fail = model.fail_cols

    same_ref = sum(1 for c in f_cols if c.depends_on)
    ext_ref = sum(1 for c in f_cols if c.ref_sheets)
    ext_names = sorted({s for c in f_cols for s in c.ref_sheets})

    # 依赖深度
    by = model.by_letter
    depth: Dict[str, int] = {}

    def dep(letter: str, stack: Set[str]) -> int:
        if letter in depth:
            return depth[letter]
        if letter in stack:
            return 99
        c = by.get(letter)
        if not c or not c.depends_on:
            depth[letter] = 0
            return 0
        stack.add(letter)
        v = 1 + max((dep(d, stack) for d in c.depends_on), default=0)
        stack.discard(letter)
        depth[letter] = v
        return v

    for c in f_cols:
        dep(c.letter, set())
    max_depth = max([v for v in depth.values() if v < 99], default=0)
    cycles = sorted({k for k, v in depth.items() if v == 99})

    unsupported_calls: Dict[str, List[str]] = {}
    for c in f_cols:
        if c.unsupported:
            unsupported_calls[c.name] = c.unsupported

    auto = len(f_cols) - len(fail)
    cov = (auto / len(f_cols) * 100) if f_cols else 100.0

    L: List[str] = []
    A = L.append
    A("# 迁移评估报告")
    A("")
    A(f"- 源文件：`{model.path}`")
    A(f"- 目标 sheet：`{model.sheet_name}`")
    A(f"- 表头行：{model.header_row}　样板行：{model.target_row}")
    A(f"- 主键列：{model.primary_key_letter or '未识别'}"
      f"{'（' + model.resolve_name(model.primary_key_letter) + '）' if model.primary_key_letter else ''}")
    A("")
    A("## 【规模】")
    A("")
    A(f"- 总列数：{len(cols)}")
    A(f"- 公式列：{len(f_cols)}")
    A(f"- 数据列：{len(d_cols)}")
    A(f"- 空列　：{len(cols) - len(f_cols) - len(d_cols)}")
    A("")
    A("## 【依赖复杂度】")
    A("")
    A(f"- 同表引用：{same_ref} 列")
    A(f"- 跨表引用：{ext_ref} 列，涉及 {len(ext_names)} 张外部表")
    A(f"- 依赖深度：最大 {max_depth} 层")
    if cycles:
        A(f"- ⚠️ 循环引用：{', '.join(cycles)}")
    A("")
    A("## 【分层情况】")
    A("")
    A(f"- Layer 1（直接引用）：{len(model.layer1)} 列"
      + (f"　→ {', '.join(c.name for c in model.layer1)}" if model.layer1 else ""))
    A(f"- Layer 2（跨表操作）：{len(model.layer2)} 列")
    A(f"- Layer 3（其他公式）：{len(model.layer3)} 列")
    A(f"- 分层冲突检测：{'✅ 0 处（使用分层模式）' if not model.conflicts else f'⚠️ {len(model.conflicts)} 处'}")
    if model.conflicts:
        A("")
        A("  > 分层模型要求「低层不依赖高层」（L1 → L2 → L3 顺序执行）。")
        A("  > 本表存在跨层反向依赖，**已自动降级为全局拓扑排序执行**（结果等价，只是不再分层）。")
        A(f"  > 降级后执行顺序由依赖图决定，不影响生成结果的正确性。")
        A("")
        for x in model.conflicts[:20]:
            A(f"    - {x}")
        if len(model.conflicts) > 20:
            A(f"    - ...（其余 {len(model.conflicts) - 20} 处见模型对象）")
    A("")
    A("## 【外部表】")
    A("")
    if not model.externals:
        A("- 无跨表引用")
    else:
        A("| 表名 | 列数 | 行数 |")
        A("|---|---|---|")
        for name in ext_names:
            es = model.externals.get(name)
            if not es:
                A(f"| {name} | - | ⚠️ 缺失 |")
                continue
            A(f"| {name} | {len(es.letters)} | {es.n_rows if es.n_rows is not None else '未扫描'} |")
    A("")
    A("## 【风险点】")
    A("")
    risks = []
    err_cols = [c for c in cols if c.cached_is_error]
    for c in err_cols:
        risks.append(f"⚠️ `{c.letter}`({c.name}) 的 Excel 缓存值是错误值 **{c.cached}**"
                     f"{'，原公式：`' + c.formula + '`' if c.formula else ''}")
    missing = [c for c in cols if c.cached_missing]
    if missing:
        risks.append(f"⚠️ {len(missing)} 个公式列没有缓存值（Excel 从未计算或未保存结果），"
                     f"对账基线缺失：{', '.join(c.name for c in missing[:10])}"
                     + (" ..." if len(missing) > 10 else ""))
    for name, fns in unsupported_calls.items():
        risks.append(f"⚠️ `{name}` 含不支持的函数：{', '.join(sorted(set(fns)))}")
    if not risks:
        risks.append("✅ 未发现明显风险点")
    for r in risks:
        A(f"- {r}")
    A("")
    A("## 【迁移成本预估】")
    A("")
    A(f"- 自动转换覆盖率：{cov:.1f}%")
    A(f"- 需人工补充列数：{len(fail)}")
    if fail:
        A("")
        A("| 列 | 列名 | 原因 |")
        A("|---|---|---|")
        for c in fail:
            A(f"| {c.letter} | {c.name} | {'、'.join(sorted(set(c.unsupported)))} |")
    A("")

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text("\n".join(L), encoding="utf-8")
    return str(out_path)

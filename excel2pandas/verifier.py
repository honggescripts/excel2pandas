# -*- coding: utf-8 -*-
"""
verifier.py —— 校验器（逐记录 + 逐字段 diff）

职责：把生成代码跑出的 df 和原始 Excel 缓存值逐格对账，输出 verify_report.md。

关键分类（真实模板逼出来的）：
  · OK        —— 一致（含容差内）
  · 精度      —— 数值差异在容差外、但在宽松容差内
  · 空值      —— 一方是 NaN
  · 逻辑      —— 值真的不一样
  · Excel侧错误 —— Excel 缓存值是 #REF! / #N/A / #NUM! 等（模板自身的问题，不算我们的逻辑差异）

不做：执行、生成。
"""
from __future__ import annotations

import math
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import numpy as np
import pandas as pd

EXCEL_ERRORS = {"#REF!", "#N/A", "#DIV/0!", "#VALUE!", "#NUM!", "#NAME?", "#NULL!", "#SPILL!", "#CALC!"}


@dataclass
class FieldDiff:
    primary_key: str
    column: str
    excel_value: Any
    python_value: Any
    diff: Optional[float]
    rel: Optional[float]
    diff_type: str          # OK / 精度 / 空值 / 逻辑 / Excel侧错误


@dataclass
class VerifyReport:
    total_columns: int = 0
    data_columns: List[str] = field(default_factory=list)
    success_columns: List[str] = field(default_factory=list)
    failed_columns: List[str] = field(default_factory=list)
    matched_rows: int = 0
    compared_cells: int = 0
    counts: Dict[str, int] = field(default_factory=dict)
    field_diffs: List[FieldDiff] = field(default_factory=list)
    excel_errors: List[Tuple[str, str, str, Any]] = field(default_factory=list)
    excel_only_columns: List[str] = field(default_factory=list)
    notes: List[str] = field(default_factory=list)

    @property
    def ok_rate(self) -> float:
        tot = self.compared_cells - self.counts.get("Excel侧错误", 0)
        return (self.counts.get("OK", 0) / tot * 100) if tot else 100.0


def is_excel_error(v: Any) -> bool:
    return isinstance(v, str) and v.strip().upper() in EXCEL_ERRORS


def _is_null(v: Any) -> bool:
    if v is None:
        return True
    if isinstance(v, float) and v != v:
        return True
    if v is pd.NaT:
        return True
    try:
        return bool(pd.isna(v))
    except (TypeError, ValueError):
        return False


def _as_num(v: Any) -> Optional[float]:
    if isinstance(v, bool):
        return float(v)
    if isinstance(v, (int, float, np.integer, np.floating)):
        f = float(v)
        return None if f != f else f
    if isinstance(v, str):
        s = v.strip().replace(",", "")
        if s.endswith("%"):
            try:
                return float(s[:-1]) / 100.0
            except ValueError:
                return None
        try:
            return float(s)
        except ValueError:
            return None
    return None


def load_excel_side(excel_path: str | Path, sheet_name: str, model,
                    header_row: int = 1) -> pd.DataFrame:
    """按位置把 Excel 缓存值对齐到模型的列名。"""
    raw = pd.read_excel(excel_path, sheet_name=sheet_name, header=header_row - 1)
    k = min(len(raw.columns), len(model.columns))
    names = [model.columns[i].name for i in range(k)]
    names += [f"__extra_{i}" for i in range(len(raw.columns) - k)]
    raw.columns = names
    return raw


def detect_excel_errors(excel_path: str | Path, sheet_name: str, header_row: int,
                        model, max_rows: int = 20000) -> Dict[str, Dict[str, str]]:
    """
    用 openpyxl 读**原始缓存值**，找出 Excel 侧本身就是错误值的单元格。

    为什么必须单独做这一步：pandas 的 openpyxl reader 会把 `#REF!` / `#N/A` / `#NUM!`
    读成 NaN，于是这些单元格会被误判成「空值」——而实际上它们是**模板自身的错误**，
    生成器按正确逻辑重算后得到的真实数值不应被判为差异。
    """
    import openpyxl
    from openpyxl.utils import get_column_letter

    wb = openpyxl.load_workbook(str(excel_path), read_only=True, data_only=True, keep_links=False)
    if sheet_name not in wb.sheetnames:
        wb.close()
        return {}
    ws = wb[sheet_name]
    ws.reset_dimensions()

    colmap = {c.index: c.name for c in model.columns}
    pk_idx = model.pick(model.primary_key_letter).index if model.primary_key_letter else None

    errs: Dict[str, Dict[str, str]] = {}
    for row in ws.iter_rows(min_row=header_row + 1, max_row=header_row + max_rows):
        pk = None
        rowerr: Dict[str, str] = {}
        for c_idx, cell in enumerate(row, start=1):
            v = getattr(cell, "value", None)
            if pk_idx and c_idx == pk_idx:
                pk = str(v) if v is not None else None
            if isinstance(v, str) and v.strip().upper() in EXCEL_ERRORS and c_idx in colmap:
                rowerr[colmap[c_idx]] = v.strip().upper()
        if pk and rowerr:
            errs.setdefault(pk, {}).update(rowerr)
    wb.close()
    return errs


def compare_cell(xv: Any, pv: Any, rtol: float, atol: float
                 ) -> Tuple[str, Optional[float], Optional[float]]:
    if is_excel_error(xv):
        return "Excel侧错误", None, None
    if _is_null(xv) and _is_null(pv):
        return "OK", 0.0, 0.0
    if _is_null(xv) or _is_null(pv):
        return "空值", None, None

    xn, pn = _as_num(xv), _as_num(pv)
    if xn is not None and pn is not None:
        d = pn - xn
        rel = abs(d) / max(abs(xn), 1e-12)
        if math.isclose(pn, xn, rel_tol=rtol, abs_tol=atol):
            return "OK", d, rel
        if math.isclose(pn, xn, rel_tol=1e-4, abs_tol=1e-4):
            return "精度", d, rel
        return "逻辑", d, rel

    xs = str(xv).strip()
    ps = str(pv).strip()
    if xs == ps:
        return "OK", None, None
    if xs.lower() == ps.lower():
        return "OK", None, None
    return "逻辑", None, None


def verify(result_df: pd.DataFrame,
           excel_path: str | Path,
           sheet_name: str,
           model,
           header_row: int = 1,
           primary_key: Optional[str] = None,
           fail_suffix: str = "_fail",
           rtol: float = 1e-6,
           atol: float = 1e-6,
           column_tolerance: Optional[Dict[str, Tuple[float, float]]] = None,
           max_detail_rows: int = 200) -> VerifyReport:
    rep = VerifyReport()
    rep.total_columns = len(model.columns)

    pk = primary_key or model.resolve_name(model.primary_key_letter) or None
    if pk is None:
        rep.notes.append("⚠️ 未识别主键列，无法按记录对账")
        return rep

    py_cols = set(result_df.columns)
    rep.failed_columns = [c.name for c in model.formula_cols if f"{c.name}{fail_suffix}" in py_cols]
    rep.success_columns = [c.name for c in model.formula_cols
                           if c.name in py_cols and f"{c.name}{fail_suffix}" not in py_cols]
    rep.data_columns = [c.name for c in model.data_cols]

    ex = load_excel_side(excel_path, sheet_name, model, header_row)
    if pk not in ex.columns:
        rep.notes.append(f"⚠️ Excel 侧找不到主键列 {pk}，无法对账")
        return rep

    keys = set(result_df[pk].astype(str))
    ex = ex[ex[pk].astype(str).isin(keys)]
    rep.matched_rows = len(ex)
    if ex.empty:
        rep.notes.append("⚠️ Excel 侧没有与结果主键匹配的行 —— 本文件只能对账「模板里已存在的样板行」")
        return rep

    tol = dict(column_tolerance or {})
    xl_errors = detect_excel_errors(excel_path, sheet_name, header_row, model)
    counts: Dict[str, int] = {}
    compare_cols: List[str] = []
    for c in model.columns:
        if c.kind == "empty" and not c.header:
            continue
        if c.name in rep.data_columns:
            continue
        if c.name in rep.success_columns:
            compare_cols.append(c.name)
        elif c.name in rep.failed_columns:
            compare_cols.append(f"{c.name}{fail_suffix}")
        else:
            rep.excel_only_columns.append(c.name)

    for _, xrow in ex.iterrows():
        key = str(xrow[pk])
        pr = result_df[result_df[pk].astype(str) == key]
        if pr.empty:
            continue
        prow = pr.iloc[0]
        for col in compare_cols:
            if col not in ex.columns:
                continue
            xv = xrow[col]
            base = col[:-len(fail_suffix)] if col.endswith(fail_suffix) else col
            pv = prow[col] if col in prow.index else None
            r_, a_ = tol.get(base, (rtol, atol))
            # Excel 侧本身就是错误值 -> 单独归类，不参与逻辑差异判定
            xerr = xl_errors.get(key, {}).get(base)
            if xerr:
                kind, diff, rel = "Excel侧错误", None, None
                xv = xerr
            else:
                kind, diff, rel = compare_cell(xv, pv, r_, a_)
            counts[kind] = counts.get(kind, 0) + 1
            rep.compared_cells += 1
            if kind != "OK":
                rep.field_diffs.append(FieldDiff(key, col, xv, pv, diff, rel, kind))
            if kind == "Excel侧错误":
                rep.excel_errors.append((key, col, str(xv), pv))

    rep.counts = counts
    return rep


def write_verify_report(rep: VerifyReport, out_path: str | Path,
                        code_path: str = "", excel_path: str = "",
                        sheet_name: str = "") -> str:
    L: List[str] = []
    A = L.append
    A("# 对账报告")
    A("")
    A(f"- 生成代码：`{code_path}`")
    A(f"- Excel 　：`{excel_path}`　sheet `{sheet_name}`")
    A(f"- 对账记录数：{rep.matched_rows}")
    A("")
    A("## 【列分类】")
    A("")
    A(f"- 总列数　：{rep.total_columns}")
    A(f"- 数据列　：{len(rep.data_columns)}（原始输入，不做对比）")
    A(f"- 成功列　：{len(rep.success_columns)}（逐字段 diff）")
    A(f"- 失败列　：{len(rep.failed_columns)}（Excel 兜底，列名带 `_fail`）")
    A("")
    A("## 【对账结果】")
    A("")
    tot = rep.compared_cells
    A(f"- 比较单元格数：{tot}")
    for k in ("OK", "精度", "空值", "逻辑", "Excel侧错误"):
        n = rep.counts.get(k, 0)
        icon = {"OK": "✅", "精度": "⚠️", "空值": "⚪", "逻辑": "❌", "Excel侧错误": "🟡"}[k]
        pct = f"{n / tot * 100:.1f}%" if tot else "-"
        A(f"- {icon} {k:<12}：{n:>5}　({pct})")
    A("")
    A(f"**一致率（排除 Excel 侧错误）：{rep.ok_rate:.2f}%**")
    A("")

    if rep.excel_errors:
        A("## 【Excel 侧错误值】（模板自身问题，不计入逻辑差异）")
        A("")
        seen: Dict[Tuple[str, str], List[Any]] = {}
        for key, col, val, pv in rep.excel_errors:
            seen.setdefault((col, val), []).append((key, pv))
        A("| 列 | Excel 值 | 记录数 | Python 重算值（示例） |")
        A("|---|---|---|---|")
        for (col, val), items in seen.items():
            A(f"| `{col}` | `{val}` | {len(items)} | {_fmt(items[0][1])} |")
        A("")
        A("> 说明：这些单元格在 Excel 里**本身就是错误值**（如 `#REF!`：引用的区域不够宽、")
        A("> `#NUM!`：财务函数无解）。生成器按正确逻辑重算后会得到真实数值 —— 因此**不应**判为「逻辑差异」。")
        A("> 如果确认是模板公式写错了，请先修 Excel 公式；生成器代码无需改动。")
        A("")

    if rep.counts.get("逻辑", 0) or rep.counts.get("精度", 0) or rep.counts.get("空值", 0):
        A("## 【差异明细】")
        A("")
        A("| 记录 | 列 | Excel 值 | Python 值 | 差值 | 相对误差 | 类型 |")
        A("|---|---|---|---|---|---|---|")
        for d in rep.field_diffs:
            if d.diff_type == "Excel侧错误":
                continue
            dv = "-" if d.diff is None else f"{d.diff:.6g}"
            rv = "-" if d.rel is None else f"{d.rel:.3e}"
            A(f"| `{d.primary_key}` | {d.column} | {_fmt(d.excel_value)} | {_fmt(d.python_value)} | {dv} | {rv} | {d.diff_type} |")
        A("")

    if rep.failed_columns:
        A("## 【失败列】")
        A("")
        for c in rep.failed_columns:
            A(f"- `{c}{'_fail'}`")
        A("")

    if rep.excel_only_columns:
        A("## 【未参与对账的列】")
        A("")
        A(f"- {len(rep.excel_only_columns)} 列：{', '.join(rep.excel_only_columns[:20])}"
          + (" …" if len(rep.excel_only_columns) > 20 else ""))
        A("")

    if rep.notes:
        A("## 【备注】")
        A("")
        for n in rep.notes:
            A(f"- {n}")
        A("")

    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text("\n".join(L), encoding="utf-8")
    return str(out)


def _fmt(v: Any) -> str:
    if v is None:
        return "-"
    if isinstance(v, float):
        if v != v:
            return "NaN"
        return f"{v:.6g}"
    if isinstance(v, (np.integer, np.floating)):
        return f"{float(v):.6g}"
    s = str(v)
    return s if len(s) <= 40 else s[:37] + "..."
